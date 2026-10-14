"""Turns the UK government food tables into data/cofid.json for the server.

McCance and Widdowson's Composition of Foods Integrated Dataset (CoFID),
the Excel file from gov.uk. Only the "Proximates" sheet is used: calories,
protein, carbohydrate and fat per 100g (per 100ml for alcoholic drinks,
group Q).

    pip install openpyxl
    python scripts/import_cofid.py path/to/CoFID.xlsx

Contains public sector information licensed under the Open Government
Licence v3.0.
"""

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


def number(v):
    """A value from the sheet: 'Tr' (trace) is 0, 'N' (not known) is None."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if s.lower() in ("tr", "trace"):
        return 0.0
    try:
        return float(s)
    except ValueError:
        return None


def find(headers, *patterns):
    for i, h in enumerate(headers):
        text = str(h or "").strip().lower()
        if any(re.search(p, text) for p in patterns):
            return i
    raise SystemExit(f"Column not found: {patterns}. Headers: {headers}")


def main(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    sheet = next((wb[n] for n in wb.sheetnames if n.strip().lower().startswith("1.3 proximates")
                  or n.strip().lower() == "proximates"), None)
    if sheet is None:
        sheet = next(wb[n] for n in wb.sheetnames if "proximates" in n.lower())
    rows = list(sheet.iter_rows(values_only=True))
    headers = rows[0]
    col = {
        "code": find(headers, r"^food code"),
        "name": find(headers, r"^food name"),
        "desc": find(headers, r"^description"),
        "group": find(headers, r"^group"),
        "kcal": find(headers, r"^energy \(kcal\)"),
        "protein": find(headers, r"^protein"),
        "fat": find(headers, r"^fat"),
        "carbs": find(headers, r"^carbohydrate"),
    }

    foods, skipped = [], 0
    # Rows 2 and 3 are codes and units; the foods start on row 4.
    for row in rows[3:]:
        name = str(row[col["name"]] or "").strip()
        if not name:
            continue
        per100 = {
            "calories": number(row[col["kcal"]]),
            "protein": number(row[col["protein"]]),
            "carbs": number(row[col["carbs"]]),
            "fat": number(row[col["fat"]]),
        }
        if any(v is None or v < 0 for v in per100.values()):
            skipped += 1
            continue
        group = str(row[col["group"]] or "").strip()
        food = {
            "code": str(row[col["code"]] or "").strip(),
            "name": name,
            "group": group,
            "per100": {k: round(v, 2) for k, v in per100.items()},
        }
        if group.upper().startswith("Q"):
            food["per100ml"] = True
        desc = str(row[col["desc"]] or "").strip()
        if desc:
            food["desc"] = desc[:160]
        foods.append(food)

    out = Path(__file__).resolve().parent.parent / "data" / "cofid.json"
    out.write_text(json.dumps({
        "_about": "McCance and Widdowson's Composition of Foods Integrated Dataset 2021 "
                  "(Proximates). Contains public sector information licensed under the "
                  "Open Government Licence v3.0. Made by scripts/import_cofid.py.",
        "foods": foods,
    }, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"{len(foods)} foods written to {out} ({skipped} skipped: values not known)")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    main(sys.argv[1])
